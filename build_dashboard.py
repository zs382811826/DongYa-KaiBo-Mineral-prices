#!/usr/bin/env python3
"""矿产行情看板生成脚本 —— 读取 Excel 生成 HTML 看板"""
import openpyxl
import json
import os
import math
from datetime import datetime, timedelta
from collections import defaultdict

EXCEL = r"D:\矿产行情看板\my_website\矿产行情记录表.xlsx"
OUTPUT = r"D:\矿产行情看板\my_website\index.html"

CATEGORIES = [
    ("稀土行情", ["氧化钕", "氧化镨钕", "氧化铽", "氧化镝", "氧化钇", "独居石"]),
    ("锆业行情", ["锆精矿(国内)", "锆精矿(越南)", "高级硅酸锆"]),
    ("钛业行情", ["钛精矿", "金红石"]),
    ("小金属行情", ["锡锭", "氧化钽", "氧化铌"]),
]

DISPLAY_NAMES = {
    "氧化钕": "氧化钕(Nd2O3/TREO：99.0%-99.9%)",
	"氧化镨钕": "氧化镨钕（Pr6O11+Nd2O3/TREO≥99%、Nd2O3/TREO≥75%）",
	"氧化铽": "氧化铽（Tb4O7/TREO：99.99%）",
	"氧化镝": "氧化镝（Dy2O3/TREO：99.5%-99.9%）",
	"氧化钇": "氧化钇（Y2O3/TREO：99.99%－99.999%）",
	"独居石": "独居石(REO=54,Pr+Nd=22.5,Tb=0.11,Dy=0.58)",
    "高级硅酸锆": "广东高级硅酸锆((Zr+Hf)O2≥64.5%,D50=1.0μm)",
    "金红石": "金红石型钛白粉（氯化法）",
    "锆精矿(国内)": "锆精矿(国内ZrO2>65%,TiO2<0.15%,Fe2O3<0.1%)",
    "锆精矿(越南)": "锆精矿(越南ZrO2>65%,TiO2<0.15%,Fe2O3<0.1%)",
    "氧化钽": "氧化钽(Ta2O5≥99% 工业级)",
    "氧化铌": "氧化铌(Nb2O5≥99.5% 冶金级)",
    "钛精矿": "钛精矿(TiO2>49-50%,Fe2O3≤8%)",
}

def get_display_name(name):
    return DISPLAY_NAMES.get(name, name)

def load_data(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sheet1"]
    headers = [c.value for c in ws[2]]
    rows = []
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
        vals = [c.value for c in row]
        if vals[0] is None or not isinstance(vals[0], datetime):
            continue
        rows.append(vals)
    return headers, rows

def fmt_wan(val):
    """Format value in wan yuan"""
    if val is None or (isinstance(val, float) and math.isnan(val)):
        return "-"
    v = val / 10000
    if v >= 10000:
        return f"{v/10000:.2f}亿"
    if v >= 1:
        return f"{v:.2f}万"
    return f"{v:.4f}万"

def fmt_pct(val):
    if val is None:
        return "-"
    sign = "+" if val > 0 else ""
    return f"{sign}{val:.2f}%"

def pct_class(val):
    if val is None:
        return ""
    return "up" if val > 0 else ("down" if val < 0 else "flat")

def compute_changes(headers, rows):
    products = headers[1:]
    if not rows:
        return products, [], [], [], {}, []

    latest = rows[-1]
    prev = rows[-2] if len(rows) >= 2 else None

    # Weekly: latest vs previous row
    weekly = []
    for i, name in enumerate(products, 1):
        if prev and prev[i] is not None and latest[i] is not None and prev[i] != 0:
            w = (latest[i] - prev[i]) / prev[i] * 100
        else:
            w = None
        weekly.append(w)

    # Monthly: average by month
    month_groups = defaultdict(list)
    for r in rows:
        d = r[0]
        key = (d.year, d.month)
        for i, name in enumerate(products, 1):
            if r[i] is not None:
                month_groups[(key, name)].append(r[i])

    month_avgs = {}
    for (key, name), vals in month_groups.items():
        month_avgs[(key, name)] = sum(vals) / len(vals)

    sorted_months = sorted(set(k[0] for k in month_groups.keys()))
    cur_month = sorted_months[-1] if sorted_months else None
    prev_month = sorted_months[-2] if len(sorted_months) >= 2 else None

    monthly = []
    for name in products:
        if cur_month and prev_month and (cur_month, name) in month_avgs and (prev_month, name) in month_avgs:
            cur_avg = month_avgs[(cur_month, name)]
            prev_avg = month_avgs[(prev_month, name)]
            m = (cur_avg - prev_avg) / prev_avg * 100 if prev_avg != 0 else None
        else:
            m = None
        monthly.append(m)

    # Yearly: latest vs earliest this year
    this_year = latest[0].year
    year_rows = [r for r in rows if r[0].year == this_year]
    earliest = year_rows[0] if year_rows else latest

    yearly = []
    for i, name in enumerate(products, 1):
        if earliest[i] is not None and latest[i] is not None and earliest[i] != 0:
            y = (latest[i] - earliest[i]) / earliest[i] * 100
        else:
            y = None
        yearly.append(y)

    # Trend data for charts
    trend_data = {}
    for i, name in enumerate(products, 1):
        points = []
        for r in rows:
            if r[i] is not None:
                points.append({"date": r[0].strftime("%Y-%m-%d"), "price": r[i]})
        trend_data[name] = points

    return products, weekly, monthly, yearly, trend_data, rows
def build_html(products, latest_prices, weekly, monthly, yearly, trend_data, rows):
    date_range = f"{rows[0][0].strftime('%Y-%m-%d')} ~ {rows[-1][0].strftime('%Y-%m-%d')}"
    total = len(products)
    avg_w = sum(w for w in weekly if w is not None) / max(1, len([w for w in weekly if w is not None]))
    avg_m = sum(m for m in monthly if m is not None) / max(1, len([m for m in monthly if m is not None]))
    avg_y = sum(y for y in yearly if y is not None) / max(1, len([y for y in yearly if y is not None]))

    idx_map = {name: i for i, name in enumerate(products)}

    rows_html = ""
    for cat_name, cat_products in CATEGORIES:
        rows_html += f'<tr class="cat-header"><td colspan="5">{cat_name}</td></tr>'
        for name in cat_products:
            i = idx_map[name]
            display = get_display_name(name)
            lp = latest_prices[i]
            w = weekly[i]
            m = monthly[i]
            y = yearly[i]
            rows_html += f"""<tr>
            <td class="product-name">{display}</td>
            <td class="price">{fmt_wan(lp)}</td>
            <td class="{pct_class(w)}">{fmt_pct(w)}</td>
            <td class="{pct_class(m)}">{fmt_pct(m)}</td>
            <td class="{pct_class(y)}">{fmt_pct(y)}</td>
        </tr>"""

    options = ""
    for cat_name, cat_products in CATEGORIES:
        options += f'<optgroup label="{cat_name}">'
        for name in cat_products:
            display = get_display_name(name)
            options += f'<option value="{name}">{display}</option>'
        options += '</optgroup>'

    html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>矿产行情看板</title>
<script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.0/dist/chart.umd.min.js"></script>
<style>
* {{ margin:0; padding:0; box-sizing:border-box; }}
body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', 'Microsoft YaHei', sans-serif; background: #f0f2f5; color: #333; }}
.header {{ background: linear-gradient(135deg, #1a3c5e, #2a5f8f); color: #fff; padding: 24px 32px; }}
.header h1 {{ font-size: 22px; margin-bottom: 4px; }}
.header .sub {{ font-size: 13px; opacity: 0.8; }}
.container {{ max-width: 1200px; margin: 0 auto; padding: 20px 24px; }}
.cards {{ display: grid; grid-template-columns: repeat(4, 1fr); gap: 16px; margin-bottom: 24px; }}
.card {{ background: #fff; border-radius: 10px; padding: 18px 20px; box-shadow: 0 1px 4px rgba(0,0,0,0.06); }}
.card .label {{ font-size: 12px; color: #888; margin-bottom: 6px; }}
.card .value {{ font-size: 26px; font-weight: 700; }}
.card .value.up {{ color: #e74c3c; }}
.card .value.down {{ color: #27ae60; }}
.section-title {{ font-size: 16px; font-weight: 600; margin: 24px 0 12px; color: #1a3c5e; }}
table {{ width: 100%; border-collapse: collapse; background: #fff; border-radius: 10px; overflow: hidden; box-shadow: 0 1px 4px rgba(0,0,0,0.06); }}
th {{ background: #f5f7fa; padding: 10px 14px; text-align: left; font-size: 12px; color: #666; font-weight: 500; }}
td {{ padding: 10px 14px; font-size: 13px; border-top: 1px solid #f0f0f0; }}
.price {{ font-weight: 600; font-variant-numeric: tabular-nums; }}
.up {{ color: #e74c3c; font-weight: 600; }}
.down {{ color: #27ae60; font-weight: 600; }}
.flat {{ color: #999; }}
.cat-header td {{ background: #e8ecf1; font-weight: 700; font-size: 13px; color: #1a3c5e; padding: 8px 14px; }}
.product-name {{ padding-left: 24px; }}
.chart-area {{ background: #fff; border-radius: 10px; padding: 20px; margin-top: 24px; box-shadow: 0 1px 4px rgba(0,0,0,0.06); }}
.chart-controls {{ display: flex; align-items: center; gap: 12px; margin-bottom: 16px; }}
.chart-controls select {{ padding: 6px 12px; border: 1px solid #d9d9d9; border-radius: 6px; font-size: 13px; }}
.chart-wrapper {{ position: relative; height: 400px; }}
.footer {{ text-align: center; padding: 20px; font-size: 11px; color: #aaa; }}
</style>
</head>
<body>
<div class="header">
    <h1>&#x1f4ca; 东亚&amp;凯博矿产 行情看板</h1>
    <div class="sub">数据范围：{date_range} &#xff5c; 共 {total} 个品种 &#xff5c; 生成时间：{datetime.now().strftime("%Y-%m-%d %H:%M")}</div>
</div>
<div class="container">
    <div class="cards">
        <div class="card"><div class="label">品种数量</div><div class="value" style="color:#1a3c5e">{total}</div></div>
        <div class="card"><div class="label">周度平均涨跌</div><div class="value {pct_class(avg_w)}">{fmt_pct(avg_w)}</div></div>
        <div class="card"><div class="label">月度平均涨跌</div><div class="value {pct_class(avg_m)}">{fmt_pct(avg_m)}</div></div>
        <div class="card"><div class="label">年度平均涨跌</div><div class="value {pct_class(avg_y)}">{fmt_pct(avg_y)}</div></div>
    </div>

    <div class="section-title">&#x1f4cb; 涨跌幅明细</div>
    <table>
        <thead><tr><th>品种</th><th>最新报价（万元）</th><th>周度涨跌</th><th>月度涨跌</th><th>年度涨跌</th></tr></thead>
        <tbody>{rows_html}</tbody>
    </table>

    <div class="chart-area">
        <div class="chart-controls">
            <label for="productSelect">&#x1f4c8; 选择品种查看趋势：</label>
            <select id="productSelect">{options}</select>
        </div>
        <div class="chart-wrapper"><canvas id="trendChart"></canvas></div>
    </div>

    <div class="footer">矿产行情看板 &#xb7; 数据来源：瑞道金属网，上海有色网，生意社</div>
</div>
<script>
var trendData = {json.dumps(trend_data, ensure_ascii=False)};
var chart = null;

function fmtWan(v) {{
    var w = v / 10000;
    if (w >= 10000) return (w/10000).toFixed(2) + '\u4ebf';
    if (w >= 1) return w.toFixed(2) + '\u4e07';
    return w.toFixed(4) + '\u4e07';
}}

function drawChart(name) {{
    if (chart) chart.destroy();
    var points = trendData[name];
    var labels = points.map(function(p) {{ return p.date; }});
    var prices = points.map(function(p) {{ return p.price; }});
    var ctx = document.getElementById('trendChart').getContext('2d');
    chart = new Chart(ctx, {{
        type: 'line',
        data: {{
            labels: labels,
            datasets: [{{
                label: name + ' (\u4e07\u5143)',
                data: prices,
                borderColor: '#e74c3c',
                backgroundColor: 'rgba(231,76,60,0.08)',
                fill: true,
                tension: 0.3,
                pointRadius: 5,
                pointHoverRadius: 7,
                pointBackgroundColor: '#e74c3c'
            }}]
        }},
        options: {{
            responsive: true,
            maintainAspectRatio: false,
            plugins: {{
                legend: {{ display: false }},
                tooltip: {{
                    callbacks: {{
                        label: function(ctx) {{ return name + ': ' + fmtWan(ctx.parsed.y); }}
                    }}
                }}
            }},
            scales: {{
                x: {{ grid: {{ display: false }} }},
                y: {{
                    ticks: {{ callback: function(v) {{ return fmtWan(v); }} }},
                    grid: {{ color: '#f0f0f0' }}
                }}
            }}
        }}
    }});
}}

document.getElementById('productSelect').addEventListener('change', function(e) {{ drawChart(e.target.value); }});
drawChart('{products[0]}');
</script>
</body>
</html>"""
    return html

def main():
    print("读取 Excel...")
    headers, rows = load_data(EXCEL)
    print(f"共 {len(rows)} 条记录, {len(headers)-1} 个品种")

    print("计算涨跌幅...")
    products, weekly, monthly, yearly, trend_data, rows2 = compute_changes(headers, rows)

    latest_prices = [rows[-1][i] if rows[-1][i] is not None else None for i in range(1, len(headers))]

    print("生成看板...")
    html = build_html(products, latest_prices, weekly, monthly, yearly, trend_data, rows2)

    os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)
    with open(OUTPUT, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"完成！看板已生成: {OUTPUT}")

if __name__ == "__main__":
    main()
